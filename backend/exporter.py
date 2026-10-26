"""Geração dos arquivos de download (PDF, Word, TXT) a partir do texto traduzido.

O estilo do PDF espelha o do frontend (frontend/style.css): manter os dois em sincronia.
"""

import html
import io
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.shared import Pt, RGBColor

from symbols import SYMBOLS, image_name, tokenize

# Pasta do frontend: o PDF busca as imagens do banco em frontend/emoji/.
FRONTEND_DIR = Path(__file__).resolve().parent.parent / "frontend"

EMOJI_FONTS = '"Segoe UI Emoji", "Apple Color Emoji", "Noto Color Emoji"'

VISUAL_CSS = f"""
@page {{ size: A4; margin: 2cm; }}
body {{
  font-family: Arial, sans-serif, {EMOJI_FONTS};
  font-size: 32px;
  line-height: 2.0;
  color: #111;
  background: #fff;
}}
h1 {{ font-size: 40px; margin: 0 0 24px; }}
p {{ margin: 0 0 32px; page-break-inside: avoid; }}
"""

ILLITERATE_CSS = f"""
@page {{ size: A4; margin: 1.5cm; }}
body {{
  font-family: Arial, sans-serif, {EMOJI_FONTS};
  line-height: 1.4;
  color: #111;
  background: #fff;
}}
h1 {{ font-size: 48px; margin: 0 0 20px; }}
.card {{
  border: 3px solid #111;
  border-radius: 16px;
  padding: 12px 24px;
  margin-bottom: 24px;
  page-break-inside: avoid;
}}
.line {{ display: flex; align-items: center; gap: 16px; margin: 6px 0; }}
.symbols {{ font-size: 40px; min-width: 45%; white-space: nowrap; flex: none; }}
.caption {{ font-size: 16px; color: #333; }}
.card-title .symbols {{ font-size: 52px; }}
.card-title .caption {{ font-size: 20px; font-weight: bold; color: #111; }}
img.emoji {{ height: 1em; width: 1em; margin-right: 0.25em; vertical-align: -0.1em; }}
img.emoji.hi {{ height: 1.5em; width: 1.5em; margin: -0.25em 0.1em -0.25em 0; vertical-align: -0.35em; }}
"""

TITLES = {
    "visual": "Sua Receita",
    "illiterate": "💊",
}


class ExportError(Exception):
    """Erro com mensagem amigável para o usuário."""


def split_blocks(text: str) -> list[list[str]]:
    """Divide o texto em blocos (um por medicamento), separados por linhas em branco."""
    blocks = re.split(r"\n\s*\n", text.strip())
    return [
        [line.strip() for line in block.splitlines() if line.strip()]
        for block in blocks
        if block.strip()
    ]


def split_line(line: str) -> tuple[str, str]:
    """Modo analfabetismo: "☀️ 💊 | 1 comprimido de manhã" -> (símbolos, legenda em texto)."""
    symbols, _, caption = line.partition("|")
    return symbols.strip(), caption.strip()


def symbols_html(text: str) -> str:
    """Troca cada símbolo do catálogo pela imagem do banco; o resto fica como texto."""
    out = []
    for piece, symbol in tokenize(text):
        if symbol is None:
            out.append(html.escape(piece))
        else:
            name = image_name(symbol)
            cls = "emoji hi" if name.startswith("hi_") else "emoji"
            out.append(
                f'<img class="{cls}" src="emoji/{name}" alt="{html.escape(SYMBOLS[symbol])}">'
            )
    return "".join(out)


def export(text: str, mode: str, fmt: str) -> tuple[bytes, str, str]:
    """Retorna (conteúdo, media type, nome do arquivo)."""
    if fmt == "txt":
        return text.encode("utf-8-sig"), "text/plain; charset=utf-8", "receita.txt"
    if fmt == "docx":
        return (
            _to_docx(text, mode),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "receita.docx",
        )
    return _to_pdf(text, mode), "application/pdf", "receita.pdf"


def _to_html(text: str, mode: str) -> str:
    esc = html.escape
    if mode == "visual":
        css = VISUAL_CSS
        body = "".join(
            f"<p>{'<br>'.join(esc(line) for line in block)}</p>" for block in split_blocks(text)
        )
    else:
        css = ILLITERATE_CSS
        cards = []
        for block in split_blocks(text):
            lines = []
            for i, line in enumerate(block):
                symbols, caption = split_line(line)
                cls = "line card-title" if i == 0 else "line"
                lines.append(
                    f'<div class="{cls}"><span class="symbols">{symbols_html(symbols)}</span>'
                    f'<span class="caption">{esc(caption)}</span></div>'
                )
            cards.append(f'<div class="card">{"".join(lines)}</div>')
        body = "".join(cards)
    title = esc(TITLES[mode]) if mode == "visual" else symbols_html(TITLES[mode])
    return (
        f'<!DOCTYPE html><html lang="pt-BR"><head><meta charset="utf-8">'
        f"<style>{css}</style></head><body><h1>{title}</h1>{body}</body></html>"
    )


def _to_pdf(text: str, mode: str) -> bytes:
    try:
        from weasyprint import HTML
    except (ImportError, OSError) as exc:
        # No Windows, o WeasyPrint precisa das bibliotecas GTK/Pango instaladas.
        raise ExportError(
            "Geração de PDF indisponível no servidor. Baixe em Word ou TXT."
        ) from exc
    return HTML(string=_to_html(text, mode), base_url=FRONTEND_DIR.as_uri() + "/").write_pdf()


def _to_docx(text: str, mode: str) -> bytes:
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.color.rgb = RGBColor(0x11, 0x11, 0x11)

    if mode == "visual":
        normal.font.size = Pt(24)  # 32px
        normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
        normal.paragraph_format.space_after = Pt(24)
        _add_title(doc, TITLES[mode], Pt(30))
        for block in split_blocks(text):
            para = doc.add_paragraph()
            for i, line in enumerate(block):
                if i:
                    para.add_run().add_break()
                para.add_run(line)
    else:
        normal.paragraph_format.space_after = Pt(4)
        _add_title(doc, TITLES[mode], Pt(36))
        for block in split_blocks(text):
            for i, line in enumerate(block):
                symbols, caption = split_line(line)
                para = doc.add_paragraph()
                para.add_run(symbols).font.size = Pt(38 if i == 0 else 30)
                if caption:
                    run = para.add_run(f"   {caption}")
                    run.font.size = Pt(15 if i == 0 else 12)
                    run.bold = i == 0
            doc.add_paragraph()

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _add_title(doc, title: str, size) -> None:
    run = doc.add_paragraph().add_run(title)
    run.bold = True
    run.font.size = size
