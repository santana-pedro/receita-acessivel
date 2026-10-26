"""Extração de conteúdo da receita, sempre em memória (nada é gravado em disco)."""

import io
from dataclasses import dataclass
from pathlib import Path

import pdfplumber
from docx import Document

# Imagens não são aceitas: apenas PDF e documentos de texto.
SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt"}

# Abaixo disso, consideramos o PDF escaneado (sem camada de texto útil).
MIN_PDF_TEXT_CHARS = 20

# Limite do campo de texto (mesmo valor do maxlength em frontend/index.html).
MAX_TYPED_CHARS = 20000


class ExtractionError(Exception):
    """Erro com mensagem amigável para o usuário."""


@dataclass
class Extracted:
    """Texto pronto para a IA, ou binário (PDF escaneado) para leitura multimodal."""

    text: str | None = None
    data: bytes | None = None
    mime_type: str | None = None

    @property
    def is_binary(self) -> bool:
        return self.data is not None


def extract(filename: str, content: bytes) -> Extracted:
    ext = Path(filename or "").suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise ExtractionError(
            "Formato não suportado. Envie um arquivo PDF, Word (.docx) ou TXT."
        )
    if not content:
        raise ExtractionError("O arquivo enviado está vazio.")

    if ext == ".txt":
        return Extracted(text=_require_text(_decode_txt(content)))
    if ext == ".docx":
        return Extracted(text=_require_text(_read_docx(content)))
    text = _read_pdf(content)
    if len(text.strip()) >= MIN_PDF_TEXT_CHARS:
        return Extracted(text=text)
    # PDF escaneado: o Gemini lê o PDF diretamente.
    return Extracted(data=content, mime_type="application/pdf")


def from_typed_text(text: str) -> Extracted:
    """Receita digitada pelo usuário no campo de texto."""
    if not text.strip():
        raise ExtractionError("Digite a receita antes de traduzir.")
    if len(text) > MAX_TYPED_CHARS:
        raise ExtractionError(
            f"O texto é muito longo (máximo de {MAX_TYPED_CHARS} caracteres)."
        )
    return Extracted(text=text.strip())


def ocr_fallback(item: Extracted) -> str:
    """OCR com Tesseract do PDF escaneado, usado quando a leitura multimodal falha."""
    try:
        import pytesseract
    except ImportError as exc:
        raise ExtractionError("OCR indisponível no servidor.") from exc

    try:
        images = _render_pdf_pages(item.data)
        text = "\n".join(pytesseract.image_to_string(img, lang="por") for img in images)
    except pytesseract.TesseractNotFoundError as exc:
        raise ExtractionError("OCR indisponível no servidor (Tesseract não instalado).") from exc
    except Exception as exc:
        raise ExtractionError(
            "Não foi possível ler o documento. Tente enviar uma versão mais nítida."
        ) from exc
    return _require_text(text)


def _decode_txt(content: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ExtractionError("Não foi possível ler o arquivo de texto.")


def _read_docx(content: bytes) -> str:
    try:
        doc = Document(io.BytesIO(content))
    except Exception as exc:
        raise ExtractionError("Arquivo Word inválido ou corrompido.") from exc
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def _read_pdf(content: bytes) -> str:
    try:
        with pdfplumber.open(io.BytesIO(content)) as pdf:
            return "\n".join(page.extract_text() or "" for page in pdf.pages)
    except Exception as exc:
        raise ExtractionError("Arquivo PDF inválido ou corrompido.") from exc


def _render_pdf_pages(content: bytes):
    with pdfplumber.open(io.BytesIO(content)) as pdf:
        return [page.to_image(resolution=300).original for page in pdf.pages]


def _require_text(text: str) -> str:
    if not text or not text.strip():
        raise ExtractionError(
            "Não encontramos texto no documento. Tente enviar outra versão do arquivo."
        )
    return text.strip()
